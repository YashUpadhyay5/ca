import requests
import time
import sys
from decimal import Decimal
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from openpyxl import load_workbook

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

def run_verification():
    print("=" * 70)
    print("  VERIFYING CLEAN ZERO-MOCK DATA STATE & 13-SHEET EXCEL ENGINE")
    print("=" * 70)

    # 1. Login
    login_res = requests.post(
        "http://127.0.0.1:8000/api/v1/auth/login",
        json={"email": "ca@mehtaca.com", "password": "AuditPassword123!"}
    ).json()
    assert "access_token" in login_res, f"Login failed: {login_res}"
    token = login_res["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("1. Authenticated successfully as:", login_res["user"]["full_name"])

    # 2. Check Clean State (Zero Mock Data)
    stmts = requests.get("http://127.0.0.1:8000/api/v1/statements/", headers=headers).json()
    clients = requests.get("http://127.0.0.1:8000/api/v1/clients/", headers=headers).json()
    txns = requests.get("http://127.0.0.1:8000/api/v1/transactions/", headers=headers).json()

    print(f"2. Initial Data Check:")
    print(f"   • Statements count:   {len(stmts)} (Expected 0)")
    print(f"   • Clients count:      {len(clients)} (Expected 0)")
    print(f"   • Transactions count: {txns.get('total', 0)} (Expected 0)")
    assert len(stmts) == 0, f"Expected 0 statements, got {len(stmts)}"
    assert len(clients) == 0, f"Expected 0 clients, got {len(clients)}"
    assert txns.get("total", 0) == 0, f"Expected 0 transactions, got {txns.get('total')}"
    print("   ✓ Clean state confirmed! Zero mock data in system.")

    # 3. Create a realistic Test Bank Statement PDF
    test_pdf_path = Path("test_bank_statement_live.pdf")
    doc = SimpleDocTemplate(str(test_pdf_path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>HDFC BANK LIMITED</b>", styles["Title"]),
        Paragraph("Customer Name: <b>Sharma Trading & Logistics LLP</b>", styles["Normal"]),
        Paragraph("Account Number: <b>50200012345678</b> | Branch: Fort, Mumbai", styles["Normal"]),
        Paragraph("Statement Period: <b>01/04/2025 to 30/04/2025</b>", styles["Normal"]),
        Spacer(1, 15)
    ]

    table_data = [
        ["Date", "Narration / Description", "Chq/Ref No", "Value Dt", "Withdrawal (Dr)", "Deposit (Cr)", "Closing Balance"],
        ["01/04/2025", "OPENING BALANCE B/F", "", "01/04/2025", "", "", "1,50,000.00"],
        ["03/04/2025", "NEFT CR-INFOSYS TECHNOLOGIES VENDOR", "NEFT881029", "03/04/2025", "", "3,50,000.00", "5,00,000.00"],
        ["07/04/2025", "ACH DR-OFFICE RENT BKC MUMBAI", "ACH11209", "07/04/2025", "75,000.00", "", "4,25,000.00"],
        ["12/04/2025", "UPI/DR/8821/AMAZON SELLER SERVICES", "UPI771029", "12/04/2025", "18,400.00", "", "4,06,600.00"],
        ["18/04/2025", "RTGS DR-ADVANCE TAX CHALLAN ITD", "TAX99102", "18/04/2025", "60,000.00", "", "3,46,600.00"],
        ["28/04/2025", "CREDIT INTEREST CAPITALIZED", "INT11029", "28/04/2025", "", "4,200.00", "3,50,800.00"]
    ]

    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.navy),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
    ]))
    story.append(t)
    doc.build(story)
    print("3. Generated realistic PDF statement: test_bank_statement_live.pdf")

    # 4. Upload the PDF statement (no client specified - auto-provisions)
    with open(test_pdf_path, "rb") as f:
        files = {"file": ("test_bank_statement_live.pdf", f, "application/pdf")}
        upload_res = requests.post(
            "http://127.0.0.1:8000/api/v1/documents/upload",
            files=files,
            headers=headers
        ).json()
    doc_id = upload_res["document_id"]
    print("4. Uploaded statement PDF. Doc ID:", doc_id)

    # 5. Trigger processing
    proc_res = requests.post(f"http://127.0.0.1:8000/api/v1/documents/{doc_id}/process", headers=headers).json()
    job_id = proc_res["job_id"]
    print("5. Triggered extraction job:", job_id)

    # 6. Poll until completed
    statement_id = None
    for _ in range(20):
        time.sleep(1)
        j_status = requests.get(f"http://127.0.0.1:8000/api/v1/documents/jobs/{job_id}", headers=headers).json()
        print(f"   Progress: {j_status['progress_percentage']}% | Step: {j_status['current_step']}")
        if j_status["status"] == "COMPLETED":
            statement_id = j_status["statement_id"]
            break

    assert statement_id, "Processing failed or timed out"
    print("6. Processing COMPLETED! Statement ID:", statement_id)

    # 7. Verify Statement & Dashboard Data
    stmt_data = requests.get(f"http://127.0.0.1:8000/api/v1/statements/{statement_id}", headers=headers).json()
    print("7. Extracted Statement Metadata:")
    print(f"   • Bank:           {stmt_data['bank_name']}")
    print(f"   • Account:        {stmt_data['account_number_detected']}")
    print(f"   • Transactions:   {stmt_data['total_transactions']}")
    print(f"   • Total Credits:  ₹{stmt_data['total_credits']}")
    print(f"   • Total Debits:   ₹{stmt_data['total_debits']}")
    print(f"   • Closing Bal:    ₹{stmt_data['closing_balance']}")
    print(f"   • Reconciliation: {stmt_data['reconciliation_status']}")

    assert stmt_data["total_transactions"] >= 5, f"Expected >=5 transactions, got {stmt_data['total_transactions']}"
    assert stmt_data["reconciliation_status"] == "RECONCILED", f"Expected RECONCILED, got {stmt_data['reconciliation_status']}"

    # 8. Test 13-Sheet Excel Export
    excel_res = requests.post(
        "http://127.0.0.1:8000/api/v1/exports/excel",
        json={"statement_id": statement_id},
        headers=headers
    )
    assert excel_res.status_code == 200, f"Excel export failed with {excel_res.status_code}"

    exported_path = Path("downloaded_analysis.xlsx")
    with open(exported_path, "wb") as f:
        f.write(excel_res.content)
    print("8. Downloaded Excel Export successfully. File size:", len(excel_res.content), "bytes")

    # 9. Verify Excel Workbook 13 Sheets, Table, and Formulas
    wb = load_workbook(str(exported_path))
    expected_13_sheets = [
        "01_Cover",
        "02_Summary",
        "03_Transactions",
        "04_Monthly_Analysis",
        "05_Debit_Analysis",
        "06_Credit_Analysis",
        "07_Payment_Mode",
        "08_Category_Analysis",
        "09_Top_Transactions",
        "10_Daily_Analysis",
        "11_Reconciliation",
        "12_Review_Required",
        "13_Reference_Data"
    ]
    print("9. Verifying Excel Worksheets:")
    for s in expected_13_sheets:
        assert s in wb.sheetnames, f"Missing sheet: {s}"
        print(f"   ✓ Sheet present: {s}")

    # Check tblTransactions
    tx_sheet = wb["03_Transactions"]
    assert "tblTransactions" in tx_sheet.tables, "Missing tblTransactions table in 03_Transactions!"
    print("   ✓ Official Excel Table 'tblTransactions' confirmed!")

    # Check Summary Formulas
    sum_sheet = wb["02_Summary"]
    assert sum_sheet["C8"].value == "=SUM(tblTransactions[Credit])", f"Unexpected formula: {sum_sheet['C8'].value}"
    assert sum_sheet["C9"].value == "=SUM(tblTransactions[Debit])", f"Unexpected formula: {sum_sheet['C9'].value}"
    print("   ✓ Structured references formula =SUM(tblTransactions[Credit]) confirmed!")

    # Cleanup temporary test files
    if test_pdf_path.exists():
        test_pdf_path.unlink()
    if exported_path.exists():
        exported_path.unlink()

    print("=" * 70)
    print("  ALL VERIFICATIONS PASSED 100%! SYSTEM IS AUDIT-READY.")
    print("=" * 70)

if __name__ == "__main__":
    run_verification()
